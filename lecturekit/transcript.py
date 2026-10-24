# -*- coding: utf-8 -*-
"""전사본 읽기·쓰기: SRT/VTT, 타임스탬프 TXT, 클로바노트 TXT·DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from .common import hms, parse_clock, srt_time

TS_RE = re.compile(r"(\d{1,2}:\d{2}(?::\d{2})?(?:[.,]\d+)?)")
TRANSCRIPT_EXTS = (".txt", ".docx", ".srt", ".vtt")


def read_text_file(path: Path) -> str:
    """.docx(클로바노트 Word 내보내기)는 문단을 줄로, 그 밖은 텍스트로 읽음"""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        import docx

        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:  # 표 안에 발화가 들어 있는 형식 대비
            for row in t.rows:
                lines += [c.text for c in row.cells]
        return "\n".join(lines)
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def _is_timestamp_line(s: str, m) -> bool:
    """타임스탬프 줄인지 판단.
    · 줄 맨 앞: '00:05 텍스트', '[00:05] 텍스트'
    · 줄 맨 끝(짧은 줄): '참석자 1 00:05', 'Attendees 2 01:02:03' (클로바노트)
    본문 속 '3:00에 시작'이나 머리말 '오후 3:00', '2026.10.07' 등은 제외"""
    if m is None:
        return False
    head, tail = s[:m.start()], s[m.end():].strip(" ])")
    if re.search(r"(오전|오후|AM|PM|am|pm)\s*$", head) or re.search(r"\d{4}\s*[./-]\s*\d{1,2}\s*[./-]\s*\d{1,2}", s):
        return False
    return m.start() <= 1 or (not tail and len(s) <= 40)


def load_transcript(path: Path, duration: float):
    """전사본을 [{start, end, text}]로. 타임스탬프가 하나도 없으면 RuntimeError."""
    raw = read_text_file(path).replace("\r", "")
    segs = []
    if "-->" in raw:
        for block in re.split(r"\n\s*\n", raw):
            lines = block.strip().split("\n")
            idx = next((i for i, l in enumerate(lines) if "-->" in l), None)
            if idx is None:
                continue
            m = re.search(r"([\d:.,]+)\s*-->\s*([\d:.,]+)", lines[idx])
            if not m:
                continue
            a, b = parse_clock(m.group(1)), parse_clock(m.group(2))
            text = re.sub(r"<[^>]+>", "", " ".join(l.strip() for l in lines[idx + 1:] if l.strip()))
            if a is not None and text:
                segs.append({"start": a, "end": b if b is not None else a, "text": text})
    else:
        # 첫 타임스탬프 이전 줄(제목·녹음 날짜 등 머리말)은 버림
        cur_t, buf = None, []
        for line in raw.split("\n"):
            s = line.strip()
            if not s or (cur_t is not None and re.fullmatch(r"#+\s.*|-{3,}", s)):  # 사람이 넣은 제목·구분선
                continue
            m = TS_RE.search(s)
            if _is_timestamp_line(s, m):
                if cur_t is not None and buf:
                    segs.append({"start": cur_t, "end": None, "text": " ".join(buf)})
                buf, cur_t = [], parse_clock(m.group(1))
                rest = s[m.end():].strip(" ])-:\t")
                if rest and m.start() <= 1:
                    buf.append(rest)
            elif cur_t is not None:
                buf.append(s)
        if cur_t is not None and buf:
            segs.append({"start": cur_t, "end": None, "text": " ".join(buf)})
        segs.sort(key=lambda x: x["start"])
        for k, sg in enumerate(segs):
            sg["end"] = segs[k + 1]["start"] if k + 1 < len(segs) else max(duration, sg["start"])
    segs.sort(key=lambda x: x["start"])
    if not segs:
        raise RuntimeError(f"전사본에서 타임스탬프를 찾지 못했습니다: {Path(path).name}\n"
                           f"  클로바노트에서 내려받을 때 '타임스탬프 포함'을 선택했는지 확인하세요.")
    return segs


def write_transcript(segs, out_dir: Path, name: str = "transcript.srt"):
    with open(Path(out_dir) / name, "w", encoding="utf-8") as f:
        for k, s in enumerate(segs, 1):
            f.write(f"{k}\n{srt_time(s['start'])} --> {srt_time(s['end'])}\n{s['text']}\n\n")


def find_external_transcript(video: Path):
    """영상과 같은 폴더에서 같은 이름의 전사본(.txt/.docx/.srt/.vtt)을 찾음.
    '영상이름_audio.txt'처럼 오디오 추출 파일 이름으로 받은 전사본도 찾음."""
    video = Path(video)
    for stem in (video.stem, f"{video.stem}_audio"):
        for ext in TRANSCRIPT_EXTS:
            p = video.with_name(stem + ext)
            if p.is_file():
                return p
    return None
